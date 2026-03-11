import os
import mimetypes
from minio import Minio
from config import MINIO_ENDPOINT, MINIO_ACCESS_KEY, MINIO_SECRET_KEY, MINIO_USE_SSL, MINIO_BUCKET
import tempfile
from utils_retry import retry_async
from langchain_core.documents import Document

from langchain_community.document_loaders import (
    PyPDFLoader,
    Docx2txtLoader,
    TextLoader,
    CSVLoader
)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt", ".md", ".csv", ".xlsx", ".xls"}

MIME_MAP = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/msword": ".doc",
    "text/plain": ".txt",
    "text/markdown": ".md",
    "text/csv": ".csv",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/vnd.ms-excel": ".xls",
}

class MinioService:
    def __init__(self):
        self.client = Minio(
            MINIO_ENDPOINT,
            access_key=MINIO_ACCESS_KEY,
            secret_key=MINIO_SECRET_KEY,
            secure=MINIO_USE_SSL,
        )
        self.bucket = MINIO_BUCKET

    @retry_async(max_retries=3, delay=1, backoff=2)
    async def load_documents(self, object_name: str, original_name: str = None):
        print(f"--> [MinIO] Đang tải file: {object_name}")
        file_extension = os.path.splitext(object_name)[1].lower()
        
        if not file_extension:
            try:
                stat = self.client.stat_object(self.bucket, object_name)
                file_extension = MIME_MAP.get(stat.content_type)
            except: pass

        if not file_extension and original_name:
            file_extension = os.path.splitext(original_name)[1].lower()

        documents = []
        temp_file_path = None

        with tempfile.NamedTemporaryFile(delete=False, suffix=file_extension) as tmp:
            temp_file_path = tmp.name
            try:
                response = self.client.get_object(self.bucket, object_name)
                for chunk in response.stream(32 * 1024):
                    tmp.write(chunk)
                response.close()
                response.release_conn()
            except Exception as e:
                if temp_file_path and os.path.exists(temp_file_path): os.remove(temp_file_path)
                raise RuntimeError(f"Lỗi tải file MinIO: {e}")

        try:
            if file_extension == ".pdf":
                loader = PyPDFLoader(temp_file_path)
                documents = loader.load()
            elif file_extension in [".docx", ".doc"]:
                loader = Docx2txtLoader(temp_file_path)
                documents = loader.load()
            elif file_extension in [".txt", ".md"]:
                loader = TextLoader(temp_file_path, encoding="utf-8")
                documents = loader.load()
            elif file_extension == ".csv":
                loader = CSVLoader(temp_file_path, encoding="utf-8")
                documents = loader.load()
            elif file_extension in [".xlsx", ".xls"]:
                # Loader Excel nhẹ bằng openpyxl (thay cho Unstructured)
                import openpyxl
                wb = openpyxl.load_workbook(temp_file_path, data_only=True)
                content = []
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        content.append(" ".join([str(cell) for cell in row if cell is not None]))
                documents = [Document(page_content="\n".join(content))]
        except Exception as e:
            print(f"--> [MinIO] Lỗi parse: {e}")
        finally:
            if temp_file_path and os.path.exists(temp_file_path):
                os.remove(temp_file_path)

        if not documents: return []

        for doc in documents:
            doc.metadata["source"] = object_name
            doc.metadata["file_name"] = original_name or object_name

        return documents